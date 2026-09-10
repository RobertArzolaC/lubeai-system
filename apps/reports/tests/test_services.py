"""Tests for the report export service."""

from io import BytesIO

from django.test import TestCase
from openpyxl import load_workbook

from apps.reports import factories, models
from apps.reports.services import ReportExportService


class ReportExportServiceTests(TestCase):
    """Tests for :class:`ReportExportService`."""

    def setUp(self) -> None:
        self.report = factories.ReportFactory()
        factories.LabAnalysisFactory(report=self.report)

    def test_preview_returns_columns_and_rows(self) -> None:
        """The preview exposes the column headers and report rows."""
        preview = ReportExportService.preview(models.Report.objects.all())
        self.assertIn("No. Lab", preview["columns"])
        self.assertEqual(len(preview["rows"]), 1)

    def test_preview_respects_limit(self) -> None:
        """The preview honours the requested row limit."""
        factories.ReportFactory.create_batch(3)
        preview = ReportExportService.preview(models.Report.objects.all(), limit=2)
        self.assertEqual(len(preview["rows"]), 2)

    def test_export_to_response_returns_buffer_and_filename(self) -> None:
        """The export returns a workbook buffer and a timestamped filename."""
        buffer, filename = ReportExportService.export_to_response(
            models.Report.objects.all()
        )
        self.assertIsInstance(buffer, BytesIO)
        self.assertTrue(filename.endswith(".xlsx"))

    def test_export_workbook_contains_data(self) -> None:
        """The generated workbook contains a data row for the report."""
        buffer, _ = ReportExportService.export_to_response(models.Report.objects.all())
        workbook = load_workbook(buffer)
        worksheet = workbook.active

        headers = [cell.value for cell in worksheet[5]]
        data_row = [cell.value for cell in worksheet[6]]
        mapping = dict(zip(headers, data_row))

        self.assertEqual(mapping["No. Lab"], self.report.lab_number)

    def test_equipment_columns_are_aligned(self) -> None:
        """The Cliente and Equipo columns map to branch and machine names."""
        buffer, _ = ReportExportService.export_to_response(models.Report.objects.all())
        worksheet = load_workbook(buffer).active

        headers = [cell.value for cell in worksheet[5]]
        data_row = [cell.value for cell in worksheet[6]]
        mapping = dict(zip(headers, data_row))

        self.assertEqual(mapping["Cliente"], self.report.machine.branch.name)
        self.assertEqual(mapping["Equipo"], self.report.machine.name)
        self.assertEqual(mapping["Modelo"], self.report.machine.model)
